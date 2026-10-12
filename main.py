"""Digilocker BGV report: pull data from MySQL and split into Lateral / Campus Excel files."""
import os
from datetime import date, datetime

import pandas as pd
import pymysql
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv(os.path.join(os.path.dirname(__file__), "jay1", ".env"))

OUT_DIR = os.path.join(os.path.dirname(__file__), "output")

QUERY = """
SELECT case_ars_no,
case_check_id,
ecf.case_flex_field6 AS 'RequestId',
ecf.case_flex_field1 AS 'CandidateId',
ecf.case_flex_field2 AS 'AssociateId',
fn_candidate_name(ecm.candidate_id) AS 'FullName',
ecf.case_flex_field10 AS 'BGVType',
ecf.case_flex_field11 AS BUName,
ecf.case_flex_field15 AS DeptName,
ecf.case_flex_field8 AS 'Account',
ecf.case_flex_field3 AS 'DOJ',
received_date AS 'BGV Initiated Date',
(SELECT created_on FROM ec_checks_auto_closure_queues ecac WHERE ecac.case_check_id=ecc.case_check_id AND queue_type=12 AND STATUS<>1 ORDER BY id DESC LIMIT 1) 'Digilocker Applicable Status',
(SELECT action_taken_on FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Comment Added' AND action_comments LIKE 'Current Queue - Candidate calling 1%%' ORDER BY action_id DESC LIMIT 1) AS '1st Calling Date',
(SELECT action_comments FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Comment Added' AND action_comments LIKE 'Current Queue - Candidate calling 1%%' ORDER BY action_id DESC LIMIT 1) AS '1st Calling Remarks',
(SELECT action_taken_on FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Check Push for Digilocker Journey' LIMIT 1) AS 'Digilocker 1st Trigger Sent Date',
(SELECT request_time FROM ec_emails_queue_educations eeqe WHERE ecc.case_check_id=eeqe.case_check_id LIMIT 0,1) AS '1st Reminder Sent Date',
(SELECT request_time FROM ec_emails_queue_educations eeqe WHERE ecc.case_check_id=eeqe.case_check_id LIMIT 1,1) AS '2nd Reminder Sent Date',
(SELECT request_time FROM ec_emails_queue_educations eeqe WHERE ecc.case_check_id=eeqe.case_check_id LIMIT 2,1) AS 'Digilocker Last Trigger Sent Date',
dqc_released_date AS 'Standard Mode of Verification Initiated Date',
(SELECT action_taken_on FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Comment Added' AND action_comments LIKE 'Current Queue - Candidate calling%%' ORDER BY action_id DESC LIMIT 1) AS 'Recent Calling Date',
(SELECT action_comments FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Comment Added' AND action_comments LIKE 'Current Queue - Candidate calling%%' ORDER BY action_id DESC LIMIT 1) AS 'Recent Calling Remarks',
(SELECT action_taken_on FROM ec_case_history ech WHERE ech.check_id=ecc.case_check_id AND action_taken='Digilocker Api Response' LIMIT 1) AS 'Digilocker Consent Submitted Date',
(SELECT (CASE STATUS
WHEN 2 THEN 'Pending with Candidate'
WHEN 4 THEN 'Digi Locker Response Recived'
ELSE 'Digilocker Response Failed' END
) FROM ec_checks_auto_closure_queues ecac WHERE ecac.case_check_id=ecc.case_check_id AND queue_type=12 AND STATUS<>1 ORDER BY id DESC LIMIT 1) AS 'Digi Locker Verification Status',
ecc.check_closure_date,
fn_vs_name(ecc.case_check_id) AS 'vs name'
FROM ec_case_master ecm
LEFT JOIN ec_case_checks ecc ON ecm.case_id=ecc.case_id
LEFT JOIN ec_case_fields ecf ON ecm.case_id=ecf.case_id
WHERE ecm.client_id=170092
AND family_id=4
AND check_created_on>=%s
AND check_created_on<=%s
AND check_status<>9
"""

START_DATE = "2025-06-30"
DATE_FMT_PY = "%d-%b-%y"
DATE_FMT_XL = "dd-mmm-yy"

DATE_COLS = [
    "DOJ", "BGV Initiated Date", "Digilocker Applicable Status", "1st Calling Date",
    "Digilocker 1st Trigger Sent Date", "1st Reminder Sent Date", "2nd Reminder Sent Date",
    "Digilocker Last Trigger Sent Date", "Standard Mode of Verification Initiated Date",
    "Recent Calling Date", "Digilocker Consent Submitted Date", "check_closure_date",
]
STATUS_COL = "Digi Locker Verification Status"
ID_COLS = ["RequestId", "CandidateId", "AssociateId"]
COMMON_HEAD = ["RequestId", "CandidateId", "AssociateId", "FullName", "BGVType", "BUName",
               "DeptName", "Account", "DOJ", "BGV Initiated Date", "Digilocker Applicable Status"]
LATERAL_COLS = COMMON_HEAD + [
    "1st Calling Date", "1st Calling Remarks", "Digilocker 1st Trigger Sent Date",
    "1st Reminder Sent Date", "2nd Reminder Sent Date", "Digilocker Last Trigger Sent Date",
    "Standard Mode of Verification Initiated Date", "Recent Calling Date", "Recent Calling Remarks",
    "Digilocker Consent Submitted Date", STATUS_COL, "check_closure_date",
]
CAMPUS_COLS = COMMON_HEAD + [
    "Digilocker 1st Trigger Sent Date", "1st Reminder Sent Date", "2nd Reminder Sent Date",
    "Digilocker Last Trigger Sent Date", "Standard Mode of Verification Initiated Date",
    "Digilocker Consent Submitted Date", STATUS_COL, "check_closure_date",
]


def fetch(start: str = START_DATE, end: str | None = None) -> pd.DataFrame:
    conn = pymysql.connect(
        host=os.environ["DB_HOST"], user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"], database=os.environ["DB_NAME"],
        port=int(os.environ.get("DB_PORT", 3306)), charset="utf8mb4",
    )
    try:
        with conn.cursor() as cur:
            cur.execute(QUERY, (start, f"{end or date.today()} 23:59:59"))
            cols = [d[0] for d in cur.description]
            return pd.DataFrame(cur.fetchall(), columns=cols)
    finally:
        conn.close()


def client_status(row) -> str:
    """Map Auth status to the client-facing status."""
    auth = str(row[STATUS_COL] or "").strip().lower()
    closed = pd.notna(row["check_closure_date"]) and str(row["check_closure_date"]).strip() != ""
    if auth.startswith("digi locker response rec"):
        return "Consent Submitted - DigiLocker Verified"
    if closed:
        return "Consent not submitted - Education Closed with Normal Method"
    return "Candidate Yet to Submit DigiLocker Consent"


def to_date(series: pd.Series) -> pd.Series:
    """Real date values (shown as dd-mmm-yy via the cell number format); unparsable text is kept as is."""
    parsed = pd.to_datetime(series, errors="coerce", dayfirst=True).dt.normalize()
    return pd.Series(
        [p.to_pydatetime() if pd.notna(p) else (None if pd.isna(o) else o) for p, o in zip(parsed, series)],
        index=series.index, dtype=object,
    )


def to_number_if_digits(series: pd.Series) -> pd.Series:
    """Digit-only IDs become numbers (no 'number stored as text' flag in Excel).
    Values with leading zeros or non-digits stay text so nothing is altered."""
    def conv(v):
        t = None if v is None else str(v).strip()
        if t and t.isascii() and t.isdigit() and not t.startswith("0"):
            return int(t)
        return t if t else None
    return series.map(conv).astype(object)


def display_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Copy with date columns as dd-mmm-yy text, for on-screen tables."""
    out = df.copy()
    for c in DATE_COLS:
        if c in out:
            out[c] = out[c].map(lambda v: v.strftime(DATE_FMT_PY) if isinstance(v, datetime) else v)
    for c in ID_COLS:
        if c in out:
            out[c] = out[c].map(lambda v: None if v is None else str(v))
    return out


def write_excel(df: pd.DataFrame, target) -> None:
    """Write df to an xlsx path/buffer: dd-mmm-yy dates, bold white-on-dark-gray header, tidy widths."""
    with pd.ExcelWriter(target, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="Report")
        ws = xw.sheets["Report"]
        header_fill = PatternFill("solid", start_color="595959", end_color="595959")
        for col_idx, name in enumerate(df.columns, start=1):
            head = ws.cell(row=1, column=col_idx)
            head.font = Font(bold=True, color="FFFFFF")
            head.fill = header_fill
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            letter = get_column_letter(col_idx)
            if name in ID_COLS:
                for cell in ws[letter][1:]:
                    if isinstance(cell.value, int):
                        cell.number_format = "0"
                    cell.alignment = Alignment(horizontal="left")
            if name in DATE_COLS:
                for cell in ws[letter][1:]:
                    if isinstance(cell.value, datetime):
                        cell.number_format = DATE_FMT_XL
                    cell.alignment = Alignment(horizontal="center")
            longest = max([len(str(name))] + [len(str(v)) for v in df[name].head(500) if v is not None])
            ws.column_dimensions[letter].width = min(max(longest + 2, 12), 45) if "Remarks" in name or name == "FullName" else min(max(longest + 2, 12), 30)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def build(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    out = df.copy()
    out[STATUS_COL] = out.apply(client_status, axis=1)  # uses raw closure date, so before date conversion
    for c in DATE_COLS:
        out[c] = to_date(out[c])
    for c in ID_COLS:
        out[c] = to_number_if_digits(out[c])
    return out[columns]


def split_reports(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Return {"Lateral": df, "Campus": df}, processed and with the agreed headers."""
    bgv = df["BGVType"].fillna("").str.strip().str.lower()
    return {
        name: build(df[bgv == name.lower()], cols)
        for name, cols in (("Lateral", LATERAL_COLS), ("Campus", CAMPUS_COLS))
    }


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    df = fetch()
    print(f"Fetched {len(df)} rows")
    bgv = df["BGVType"].fillna("").str.strip().str.lower()
    stamp = date.today().strftime("%d%b%Y")
    for name, part in split_reports(df).items():
        path = os.path.join(OUT_DIR, f"Digilocker_{name}_{stamp}.xlsx")
        write_excel(part, path)
        print(f"{name}: {len(part)} rows -> {path}")
    other = df[~bgv.isin(["lateral", "campus"])]
    if len(other):
        print(f"Skipped {len(other)} rows with other BGVType: {other['BGVType'].value_counts(dropna=False).to_dict()}")


if __name__ == "__main__":
    main()
